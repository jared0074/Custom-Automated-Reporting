import threading
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Border, PatternFill, Side
from openpyxl.utils import get_column_letter


# ---------- Helper ----------
def parse_percentage(val):
    """Convert '18.88%' → 0.1888, or leave plain floats unchanged."""
    if isinstance(val, str) and val.strip().endswith("%"):
        try:
            return float(val.strip().rstrip("%")) / 100
        except ValueError:
            return 0.0
    try:
        return float(val)
    except (TypeError, ValueError):
        return 0.0


# ---------- Report Logic ----------
def generate_report(input_file, product_file, output_file, platform, log):
    try:
        log(f"Platform: {platform}")

        # Read product file
        product_df = pd.read_excel(product_file)

        # Rename product columns
        product_df = product_df.rename(columns={
            "Product Code": "Seller SKU",
            "Product Name & Color": "Product Name",
            "Shopee Platform Charges": "Shopee Charges",
            "Tiktok Platform Charges": "Tiktok Charges",
        })

        # Normalize product SKU
        product_df["Seller SKU"] = product_df["Seller SKU"].astype(str).str.strip().str.upper()

        # ---------- Platform-specific ingestion ----------
        if platform == "tiktok":
            df = pd.read_excel(input_file, header=0, skiprows=[1])
            df["Seller SKU"] = df["Seller SKU"].astype(str).str.strip().str.upper()
            df["Adjusted Subtotal"] = (
                df["SKU Subtotal After Discount"].fillna(0)
                + df["SKU Platform Discount"].fillna(0)
            )
            df["Per Unit Adjusted Subtotal"] = (
                df["Adjusted Subtotal"] / df["Quantity"]
            ).round(5)
            df = df[(df["Adjusted Subtotal"] != 0) & (df["Quantity"] != 0)]

        else:  # shopee
            df = pd.read_excel(input_file, header=0)
            df = df.rename(columns={"SKU Reference No.": "Seller SKU"})
            df["Seller SKU"] = df["Seller SKU"].astype(str).str.strip().str.upper()
            df["Per Unit Adjusted Subtotal"] = pd.to_numeric(
                df["Deal Price"], errors="coerce"
            ).round(5)
            df = df[
                df["Per Unit Adjusted Subtotal"].notna()
                & (df["Per Unit Adjusted Subtotal"] != 0)
                & (df["Quantity"] != 0)
            ]

        # ---------- Group (shared for both platforms) ----------
        grouped = (
            df.groupby(["Seller SKU", "Per Unit Adjusted Subtotal"])
            .agg(Total_Quantity=("Quantity", "sum"))
            .reset_index()
        )
        grouped["Total_Amount"] = (
            grouped["Per Unit Adjusted Subtotal"] * grouped["Total_Quantity"]
        ).round(2)

        # Merge product info
        grouped = grouped.merge(product_df, on="Seller SKU", how="left")

        # Parse percentage strings → float rates
        grouped["Shopee Charges"] = (
            grouped["Shopee Charges"].fillna(0).apply(parse_percentage)
        )
        grouped["Tiktok Charges"] = (
            grouped["Tiktok Charges"].fillna(0).apply(parse_percentage)
        )

        # Select the right rate for the platform
        if platform == "tiktok":
            grouped["Platform Rate"] = grouped["Tiktok Charges"]
        elif platform == "shopee":
            grouped["Platform Rate"] = grouped["Shopee Charges"]
        else:
            grouped["Platform Rate"] = 0

        # Cost & profit  (platform charge = rate × total_amount, not rate × quantity)
        grouped["Cost"] = grouped["Cost"].fillna(0)
        grouped["Total_Cost"] = (grouped["Cost"] * grouped["Total_Quantity"]).round(2)
        grouped["Platform Charges"] = (
            grouped["Platform Rate"] * grouped["Total_Amount"]
        ).round(2)
        grouped["Profit"] = (
            grouped["Total_Amount"] - grouped["Total_Cost"] - grouped["Platform Charges"]
        ).round(2)

        # Build final report rows
        rows = []
        for sku, sku_group in grouped.groupby("Seller SKU"):
            sku_group = sku_group.sort_values("Per Unit Adjusted Subtotal")
            product_name = (
                sku_group["Product Name"].iloc[0]
                if "Product Name" in sku_group.columns
                else ""
            )

            sku_group_copy = sku_group.copy()
            sku_group_copy["Seller SKU"] = [sku] + [""] * (len(sku_group_copy) - 1)
            sku_group_copy["Product Name"] = [product_name] + [""] * (len(sku_group_copy) - 1)
            sku_group_copy["row_type"] = "line"
            rows.append(sku_group_copy)

            sku_total_row = pd.DataFrame({
                "Seller SKU": [f"{sku} Total"],
                "Product Name": [product_name],
                "Per Unit Adjusted Subtotal": [""],
                "Total_Quantity": [sku_group["Total_Quantity"].sum()],
                "Total_Amount": [sku_group["Total_Amount"].sum()],
                "Total_Cost": [sku_group["Total_Cost"].sum()],
                "Platform Charges": [sku_group["Platform Charges"].sum()],
                "Profit": [sku_group["Profit"].sum()],
                "row_type": ["total"],
            })
            rows.append(sku_total_row)

            spacer_row = pd.DataFrame({
                "Seller SKU": [""], "Product Name": [""],
                "Per Unit Adjusted Subtotal": [""], "Total_Quantity": [""],
                "Total_Amount": [""], "Total_Cost": [""],
                "Platform Charges": [""], "Profit": [""],
                "row_type": ["spacer"],
            })
            rows.append(spacer_row)

        # Combine
        final_df = pd.concat(rows, ignore_index=True)

        # Grand Total
        grand_row = pd.DataFrame({
            "Seller SKU": ["Grand Total"],
            "Product Name": [""],
            "Per Unit Adjusted Subtotal": [""],
            "Total_Quantity": [grouped["Total_Quantity"].sum()],
            "Total_Amount": [grouped["Total_Amount"].sum()],
            "Total_Cost": [grouped["Total_Cost"].sum()],
            "Platform Charges": [grouped["Platform Charges"].sum()],
            "Profit": [grouped["Profit"].sum()],
            "row_type": ["total"],
        })
        final_df = pd.concat([final_df, grand_row], ignore_index=True)

        # Export
        export_df = final_df.drop(columns=["row_type"])
        export_df = export_df[[
            "Seller SKU", "Product Name", "Per Unit Adjusted Subtotal",
            "Total_Quantity", "Total_Amount", "Total_Cost",
            "Platform Charges", "Profit",
        ]]
        export_df.to_excel(output_file, index=False)

        # Styling
        wb = load_workbook(output_file)
        ws = wb.active
        thin = Side(border_style="thin", color="000000")
        border = Border(top=thin, left=thin, right=thin, bottom=thin)
        yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

        for idx, row in final_df.iterrows():
            excel_row = idx + 2
            if row["row_type"] in ["line", "total"]:
                for col in range(1, 9):
                    ws[f"{get_column_letter(col)}{excel_row}"].border = border
                if row["row_type"] == "total":
                    for col in range(1, 9):
                        ws[f"{get_column_letter(col)}{excel_row}"].fill = yellow_fill

        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            if col_letter == "A":
                ws.column_dimensions[col_letter].width = 22
            elif col_letter == "B":
                ws.column_dimensions[col_letter].width = 30
            else:
                ws.column_dimensions[col_letter].width = 16

        wb.save(output_file)
        log(f"Report saved to: {output_file}")
        return True

    except Exception as exc:
        log(f"Error: {exc}")
        return False


# ---------- UI ----------
class ReportApp:
    def __init__(self, root):
        self.root = root
        root.title("KL Sales Report Generator")
        root.resizable(False, False)

        pad = {"padx": 10, "pady": 6}

        # Input file row
        tk.Label(root, text="Input File:").grid(row=0, column=0, sticky="w", **pad)
        self.input_var = tk.StringVar()
        tk.Entry(root, textvariable=self.input_var, width=55).grid(row=0, column=1, **pad)
        tk.Button(root, text="Browse…", command=self._browse_input).grid(row=0, column=2, **pad)

        # Product file row
        tk.Label(root, text="Product File:").grid(row=1, column=0, sticky="w", **pad)
        self.product_var = tk.StringVar()
        tk.Entry(root, textvariable=self.product_var, width=55).grid(row=1, column=1, **pad)
        tk.Button(root, text="Browse…", command=self._browse_product).grid(row=1, column=2, **pad)

        # Output file name row
        tk.Label(root, text="Output File Name:").grid(row=2, column=0, sticky="w", **pad)
        self.output_var = tk.StringVar(value="sales_report.xlsx")
        tk.Entry(root, textvariable=self.output_var, width=55).grid(row=2, column=1, **pad)

        # Platform row
        tk.Label(root, text="Platform:").grid(row=3, column=0, sticky="w", **pad)
        self.platform_var = tk.StringVar(value="TikTok")
        platform_cb = ttk.Combobox(
            root, textvariable=self.platform_var,
            values=["TikTok", "Shopee"], state="readonly", width=15,
        )
        platform_cb.grid(row=3, column=1, sticky="w", **pad)

        # Generate button (tk.Label used so bg/fg colour renders on macOS too)
        self.run_btn = tk.Label(
            root, text="Generate Report",
            bg="#4CAF50", fg="white", font=("", 10, "bold"),
            width=20, pady=6, cursor="hand2", relief="flat",
        )
        self.run_btn.grid(row=4, column=0, columnspan=3, pady=12)
        self.run_btn.bind("<Button-1>", lambda _: self._run())
        self.run_btn.bind("<Enter>", lambda _: self.run_btn.configure(bg="#45a049"))
        self.run_btn.bind("<Leave>", lambda _: self.run_btn.configure(bg="#4CAF50"))

        # Log area
        tk.Label(root, text="Log:").grid(row=5, column=0, sticky="nw", padx=10)
        self.log_box = tk.Text(root, height=10, width=70, state="disabled", bg="#f5f5f5")
        self.log_box.grid(row=6, column=0, columnspan=3, padx=10, pady=(0, 10))

    # ---- file browsers ----
    def _browse_input(self):
        path = filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx *.xls")])
        if path:
            self.input_var.set(path)

    def _browse_product(self):
        path = filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx *.xls")])
        if path:
            self.product_var.set(path)

    # ---- logging (safe from any thread via root.after) ----
    def _log(self, msg):
        self.root.after(0, self._append_log, msg)

    def _append_log(self, msg):
        self.log_box.configure(state="normal")
        self.log_box.insert(tk.END, msg + "\n")
        self.log_box.see(tk.END)
        self.log_box.configure(state="disabled")

    def _enable_btn(self):
        self._running = False
        self.run_btn.configure(bg="#4CAF50", fg="white", cursor="hand2")

    # ---- run ----
    def _run(self):
        if getattr(self, "_running", False):
            return

        input_file = self.input_var.get().strip()
        product_file = self.product_var.get().strip()
        output_file = self.output_var.get().strip()
        platform = self.platform_var.get().lower()  # "tiktok" or "shopee"

        if not input_file or not product_file or not output_file:
            messagebox.showerror("Missing Input", "Please fill in all file fields.")
            return

        if not output_file.lower().endswith(".xlsx"):
            output_file += ".xlsx"
            self.output_var.set(output_file)

        self._running = True
        self.run_btn.configure(bg="#aaaaaa", fg="#666666", cursor="arrow")
        self._append_log("Starting report generation…")

        def task():
            success = generate_report(input_file, product_file, output_file, platform, self._log)
            if success:
                self.root.after(0, messagebox.showinfo, "Done",
                                f"Report saved to:\n{output_file}")
            else:
                self.root.after(0, messagebox.showerror, "Failed",
                                "Report generation failed. Check the log for details.")
            self.root.after(0, self._enable_btn)

        threading.Thread(target=task, daemon=True).start()


if __name__ == "__main__":
    root = tk.Tk()
    ReportApp(root)
    root.mainloop()
